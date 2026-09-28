import os, io, json, math, re
from http.server import BaseHTTPRequestHandler, HTTPServer
from email.parser import BytesParser
from email.policy import default
from urllib.parse import urlparse
from openpyxl import Workbook, load_workbook

ALIASES = {
    "id": ["bale id","bale_id","baleid","id"],
    "origin": ["origin","location","source","region"],
    "mic": ["micronaire","mic"],
    "length": ["fiber length (mm)","fiber length","length","uhml","staple length"],
    "strength": ["strength (g/tex)","strength","tenacity"],
    "uniformity": ["uniformity (%)","uniformity","ui"],
    "rd": ["rd","rd value"],
    "pb": ["+b","b","plus b","plusb"],
    "moisture": ["moisture (%)","moisture"],
    "weight": ["weight (kg)","weight","bale weight","quantity (kg)"],
    "price": ["purchase price (₹/kg)","purchase price","price","rate","cost"],
    "date": ["purchase date","date"]
}

def norm(v):
    return re.sub(r"[^a-z0-9]+"," ", str(v or "").lower()).strip()

def num(v):
    try: return float(str(v).replace(",","").replace("₹","").strip())
    except: return None

def parse_bales(data):
    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.values)
    if not rows: return []
    headers = [norm(x) for x in rows[0]]
    idx = {}
    for key, aliases in ALIASES.items():
        for a in aliases:
            a = norm(a)
            if a in headers:
                idx[key] = headers.index(a); break
    out=[]
    for row in rows[1:]:
        if not any(x is not None for x in row): continue
        b={}
        for key in ["id","origin","date"]:
            i=idx.get(key); b[key]=str(row[i] if i is not None and i < len(row) else "") 
        if not b["id"].strip(): continue
        b["origin"] = b["origin"].strip() or "Unknown"
        for key in ["mic","length","strength","uniformity","rd","pb","moisture","weight","price"]:
            i=idx.get(key); b[key]=num(row[i]) if i is not None and i < len(row) else None
        b["weight"] = b["weight"] or 200
        out.append(b)
    return out

def parse_history(data):
    wb=load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    ws=wb[wb.sheetnames[0]]
    rows=list(ws.values)
    if not rows: return []
    h=[norm(x) for x in rows[0]]
    def col(name):
        aliases=ALIASES.get(name,[name])
        for a in aliases:
            a=norm(a)
            if a in h: return h.index(a)
        return -1
    def val(row,name):
        i=col(name); return row[i] if i>=0 and i<len(row) else None
    out=[]
    for r in rows[1:]:
        if not any(x is not None for x in r): continue
        out.append({
            "id": str(val(r,"id") or val(r,"date") or f"H-{len(out)+1}"),
            "count": str(val(r,"origin") or "40s"),
            "strength": num(val(r,"strength")),"length":num(val(r,"length")),
            "mic":num(val(r,"mic")),"uniformity":num(val(r,"uniformity")),
            "rd":num(val(r,"rd")),"pb":num(val(r,"pb")),
            "grade": str(val(r,"weight") or "A")
        })
    return out

def targets(form):
    return {
        "quantity": num(form.get("quantity")) or 10000,
        "mic_min": num(form.get("mic_min")) or 3.8,
        "mic_max": num(form.get("mic_max")) or 4.2,
        "mic_target": num(form.get("mic_target")) or 4.0,
        "length_min": num(form.get("length_min")) or 28,
        "strength_min": num(form.get("strength_min")) or 30,
        "uniformity_min": num(form.get("uniformity_min")) or 82,
        "rd_min": num(form.get("rd_min")) or 74,
        "pb_min": num(form.get("pb_min")) or 7
    }

def score_bale(b,t):
    specs={"mic":(t["mic_min"],t["mic_max"]),"length":(t["length_min"],999),
           "strength":(t["strength_min"],999),"uniformity":(t["uniformity_min"],999),
           "rd":(t["rd_min"],999),"pb":(t["pb_min"],999)}
    weights={"mic":.18,"length":.18,"strength":.24,"uniformity":.16,"rd":.12,"pb":.12}
    total=0; reasons=[]
    for k,w in weights.items():
        v=b.get(k); lo,hi=specs[k]
        if v is None: total += w*.70; continue
        if k=="mic":
            if v<lo: total += w*max(.15,1-(lo-v)/.5); reasons.append("mic below target")
            elif v>hi: total += w*max(.15,1-(v-hi)/.5); reasons.append("mic above target")
            else: total += w*max(.75,1-abs(v-t["mic_target"])/.4)
        else:
            if v<lo:
                total += w*max(.15,1-(lo-v)/max(abs(lo),1)); reasons.append(k+" below target")
            else: total += w
    return round(max(0,min(100,total*100)),1), reasons

def recommend(bales,t):
    ranked=[]
    for b in bales:
        s,r=score_bale(b,t); x=dict(b,score=s,reasons=r); ranked.append(x)
    ranked.sort(key=lambda x:-x["score"])
    selected=[]; qty=0; origins={}
    for b in ranked:
        if qty>=t["quantity"]: break
        origin=b["origin"]
        if origins.get(origin,0)>=t["quantity"]*.55: continue
        selected.append(b); qty += b["weight"] or 0; origins[origin]=origins.get(origin,0)+(b["weight"] or 0)
    if qty<t["quantity"]:
        for b in ranked:
            if b in selected: continue
            selected.append(b); qty += b["weight"] or 0
            if qty>=t["quantity"]: break
    expected={}
    for k in ["strength","length","mic","uniformity","rd","pb"]:
        vals=[b for b in selected if b.get(k) is not None]
        sw=sum((b.get("weight") or 1) for b in vals) or 1
        expected[k]=round(sum(b[k]*(b.get("weight") or 1) for b in vals)/sw,2)
    conf=round(sum(b["score"] for b in selected)/len(selected),1) if selected else 0
    return ranked,selected,qty,expected,conf

def history_validate(expected,hist):
    if not hist: return {"historical_count":0,"closest":None,"similarity":0,"variances":{}}
    def distance(h):
        vals={"strength":(h.get("strength"),expected.get("strength")),
              "length":(h.get("length"),expected.get("length")),
              "mic":(h.get("mic"),expected.get("mic")),
              "uniformity":(h.get("uniformity"),expected.get("uniformity")),
              "rd":(h.get("rd"),expected.get("rd")),
              "pb":(h.get("pb"),expected.get("pb"))}
        ds=[]
        for a,b in vals.values():
            if a is not None and b is not None: ds.append(abs(a-b)/max(abs(b),1))
        return sum(ds)/len(ds) if ds else 1
    closest=sorted(hist,key=distance)[0]
    sim=round(max(0,100-distance(closest)*100),1)
    keys=["strength","length","mic","uniformity","rd","pb"]
    variances={k:round((closest.get(k) or 0)-(expected.get(k) or 0),2) for k in keys}
    return {"historical_count":len(hist),"closest":closest,"similarity":sim,"variances":variances}

def sample_bales():
    wb=Workbook(); ws=wb.active; ws.title="Bales"
    ws.append(["Bale ID","Origin","Strength (g/tex)","Fiber Length (mm)","Micronaire","Uniformity (%)","Rd","+b","Weight (kg)"])
    origins=["Gujarat","Maharashtra","Telangana","Madhya Pradesh","Punjab"]
    for i in range(1,151):
        o=origins[(i-1)%len(origins)]
        ws.append([f"B{i:03d}",o,round(29+(i%13)*.55,2),round(27.2+(i%11)*.22,2),
                   round(3.65+(i%8)*.09,2),round(80.5+(i%10)*.65,2),
                   round(72.5+(i%9)*.45,2),round(6.3+(i%9)*.18,2),200])
    bio=io.BytesIO(); wb.save(bio); return bio.getvalue()

def sample_history():
    wb=Workbook(); ws=wb.active; ws.title="History"
    ws.append(["Mix ID","Count","Actual Strength","Actual Length (mm)","Actual Micronaire","Actual Uniformity (%)","Actual Rd","Actual +b","Actual Output Grade"])
    for i in range(1,31):
        ws.append([f"MIX-{i:03d}","40s",round(31+(i%6)*.35,2),round(29+(i%5)*.18,2),
                   round(3.9+(i%5)*.04,2),round(83+(i%6)*.45,2),
                   round(74.5+(i%5)*.3,2),round(7.1+(i%5)*.12,2),"A" if i%4 else "B"])
    bio=io.BytesIO(); wb.save(bio); return bio.getvalue()

HTML = open(os.path.join(os.path.dirname(__file__),"templates","index.html"),"r",encoding="utf-8").read()

class Handler(BaseHTTPRequestHandler):
    def send_bytes(self,data,ctype="application/octet-stream",status=200):
        self.send_response(status); self.send_header("Content-Type",ctype); self.send_header("Content-Length",str(len(data))); self.end_headers(); self.wfile.write(data)
    def json(self,obj,status=200): self.send_bytes(json.dumps(obj).encode(), "application/json", status)
    def do_GET(self):
        p=urlparse(self.path).path
        if p=="/": return self.send_bytes(HTML.encode(),"text/html; charset=utf-8")
        if p=="/sample/cotton_bales.xlsx": return self.send_bytes(sample_bales(),"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        if p=="/sample/historical_mixes.xlsx": return self.send_bytes(sample_history(),"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        if p=="/health": return self.json({"ok":True,"service":"CottonMix AI POC"})
        self.send_bytes(b"Not found","text/plain",404)
    def multipart(self):
        ctype=self.headers.get("Content-Type","")
        raw=self.rfile.read(int(self.headers.get("Content-Length","0")))
        msg=BytesParser(policy=default).parsebytes(b"Content-Type: "+ctype.encode()+b"\r\nMIME-Version: 1.0\r\n\r\n"+raw)
        fields={}; files={}
        for part in msg.iter_parts():
            name=part.get_param("name",header="content-disposition")
            filename=part.get_filename()
            data=part.get_payload(decode=True) or b""
            if filename: files[name]=(filename,data)
            else: fields[name]=data.decode(errors="ignore")
        return fields,files
    def do_POST(self):
        p=urlparse(self.path).path
        try:
            fields,files=self.multipart()
            if p=="/api/analyze":
                if "file" not in files: return self.json({"error":"Excel file is required"},400)
                bales=parse_bales(files["file"][1])
                if not bales: return self.json({"error":"No bale rows found in the Excel file"},400)
                ranked,selected,qty,expected,conf=recommend(bales,targets(fields))
                return self.json({"total_bales":len(bales),"selected_count":len(selected),"selected_qty":round(qty,1),
                                  "confidence":conf,"expected":expected,
                                  "selected":selected,"ranked":ranked[:25],"target":targets(fields)})
            if p=="/api/validate":
                if "historical" not in files or "bales" not in files: return self.json({"error":"Both bale and historical Excel files are required"},400)
                bales=parse_bales(files["bales"][1]); hist=parse_history(files["historical"][1])
                _,sel,qty,expected,conf=recommend(bales,targets(fields))
                v=history_validate(expected,hist)
                return self.json({"current":expected,**v})
            self.json({"error":"Not found"},404)
        except Exception as e:
            self.json({"error":str(e)},500)

if __name__=="__main__":
    port=int(os.environ.get("PORT","8080"))
    HTTPServer(("0.0.0.0",port),Handler).serve_forever()
